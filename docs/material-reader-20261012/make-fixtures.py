from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR,MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from openpyxl import Workbook
from docx import Document
import io,zipfile,json
out=Path(__file__).parent/'fixtures';out.mkdir(exist_ok=True)
img=Image.new('RGB',(700,300),'white');d=ImageDraw.Draw(img);d.text((40,50),'SCANNED SOURCE: A = 35, B = 12',fill='black',font_size=26);d.rectangle((40,110,310,200),outline='navy',width=4);d.line((310,155,600,155),fill='navy',width=4)
b=io.BytesIO();img.save(b,format='PNG');out.joinpath('scan.png').write_bytes(b.getvalue())
c=canvas.Canvas(str(out/'mixed-research.pdf'),pagesize=(612,792))
c.setTitle('Synthetic material reader QA, not research evidence')
c.drawString(50,735,'Ordinary text: Alpha 35; Beta 12. Units: samples.');c.drawString(50,705,'No significance or additional conclusions are supplied.');c.showPage()
for col,x in [('Left',45),('Right',325)]:
 for i in range(12):c.drawString(x,740-i*25,f'{col} column paragraph {i+1}: value {35 if col=="Left" else 12}')
c.showPage();c.drawString(50,740,'VECTOR FIGURE: Input -> Check -> Output');
for x,label in [(50,'Input'),(230,'Check'),(420,'Output')]:c.rect(x,580,120,70);c.drawString(x+20,610,label)
for x in [170,350]:c.line(x,615,x+60,615);c.line(x+60,615,x+50,620);c.line(x+60,615,x+50,610)
c.showPage();c.drawString(50,735,'Complex table: merged group header, exact saved values');c.rect(50,680,480,35);c.drawString(60,693,'Merged header: observations (original values, samples)')
for row,values in enumerate([['Group','Measure A','Measure B'],['Control','35','12'],['Treatment','21','9']]):
 for col,value in enumerate(values):c.rect(50+col*160,640-row*40,160,40);c.drawString(60+col*160,655-row*40,value)
c.showPage();c.drawImage(ImageReader(io.BytesIO(b.getvalue())),30,400,width=550,height=235);c.showPage();c.drawString(50,730,'Mixed page: text layer above scanned diagram');c.drawImage(ImageReader(io.BytesIO(b.getvalue())),30,400,width=550,height=235);c.save()
c=canvas.Canvas(str(out/'encrypted.pdf'),encrypt='qa-password');c.drawString(10,50,'Protected synthetic file');c.save();out.joinpath('corrupt.pdf').write_bytes(b'%PDF-1.7\ncorrupt')
p=Presentation();p.slide_width=Inches(10);p.slide_height=Inches(6)
s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_textbox(Inches(.5),Inches(.3),Inches(9),Inches(.7)).text_frame;t.text='研究材料：Alpha 35，Beta 12';t.paragraphs[0].runs[0].font.name='Noto Sans CJK SC';t.paragraphs[0].runs[0].font.size=Pt(28)
for x,label in [(1,'输入'),(4,'检查'),(7,'输出')]:
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2),Inches(2),Inches(1));sh.text=label
for x in [3,6]:
 connector=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(2.5),Inches(x+1),Inches(2.5));tip=OxmlElement('a:tailEnd');tip.set('type','triangle');connector._element.spPr.get_or_add_ln().append(tip)
s.notes_slide.notes_text_frame.text='备注专属：这是合成测试材料，不是研究结论。样本单位为个。'
s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_picture(str(out/'scan.png'),Inches(.5),Inches(.5),width=Inches(4));table=s.shapes.add_table(3,2,Inches(5),Inches(.5),Inches(4),Inches(2)).table
for i,vs in enumerate([['名称','原值'],['Alpha','35'],['Beta','12']]):
 for j,v in enumerate(vs):table.cell(i,j).text=v
chart=CategoryChartData();chart.categories=['Alpha','Beta'];chart.add_series('Observed samples',[35,12]);s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(3),Inches(8),Inches(2.6),chart)
s.notes_slide.notes_text_frame.text='图表使用原始保存数值 35 和 12，不含误差线或显著性。'
p.save(str(out/'research-slides.pptx'))
(out/'observations.csv').write_text('label,value,unit\nAlpha,35,samples\nBeta,12,samples\n')
(out/'legacy.ppt').write_bytes(bytes.fromhex('D0CF11E0A1B11AE1'))
w=Workbook();w.active.title='原始数据';w.active.append(['label','value','formula']);w.active.append(['Alpha',35,'=B2*2']);w.active.append(['Beta',12]);w.save(out/'observations.xlsx')
doc=Document();doc.add_heading('Original data',0);doc.add_paragraph('Alpha = 35; Beta = 12 samples.');table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Alpha';table.cell(0,1).text='35';table.cell(1,0).text='Beta';table.cell(1,1).text='12';doc.save(out/'observations.docx')
print(out)
# Fragmented text positions must not make an otherwise small page unreadable.
c=canvas.Canvas(str(out/'dense.pdf'),pagesize=(612,792))
for row in range(60):
 for col in range(24):c.drawString(30+col*22,760-row*12,chr(65+col%26))
c.save()
